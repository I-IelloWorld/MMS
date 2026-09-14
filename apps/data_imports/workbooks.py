from io import BytesIO

from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

from apps.assets.models import AssetCategory, Equipment
from apps.facilities.access import accessible_warehouses


EQUIPMENT_IMPORT_HEADERS = (
    "warehouse_code",
    "asset_code",
    "name",
    "category_code",
    "manufacturer",
    "model",
    "serial_number",
    "commissioned_on",
    "location_detail",
    "criticality",
    "status",
    "component_code",
    "component_name",
    "component_type",
)
TEMPLATE_LAST_ROW = 501


def _add_defined_name(workbook, name, column, count):
    last_row = max(2, count + 1)
    workbook.defined_names.add(
        DefinedName(name, attr_text=f"'选项'!${column}$2:${column}${last_row}")
    )


def _add_list_validation(sheet, cell_range, defined_name, prompt):
    validation = DataValidation(
        type="list",
        formula1=f"={defined_name}",
        allow_blank=False,
    )
    validation.error = _("请选择模板下拉列表中的有效值。")
    validation.errorTitle = _("无效选项")
    validation.prompt = prompt
    validation.promptTitle = _("请选择")
    validation.showErrorMessage = True
    validation.showInputMessage = True
    validation.showDropDown = False
    sheet.add_data_validation(validation)
    validation.add(cell_range)


def build_equipment_import_template(user):
    warehouses = list(
        accessible_warehouses(user)
        .filter(status="ACTIVE")
        .values_list("code", "name")
    )
    categories = list(
        AssetCategory.objects.filter(is_active=True)
        .order_by("code")
        .values_list("code", "name")
    )
    criticalities = [(value, str(label)) for value, label in Equipment.Criticality.choices]
    statuses = [(value, str(label)) for value, label in Equipment.Status.choices]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _("设备导入")
    sheet.sheet_view.showGridLines = False
    sheet.append(EQUIPMENT_IMPORT_HEADERS)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:N{TEMPLATE_LAST_ROW}"

    header = sheet[1]
    for cell in header:
        cell.fill = PatternFill("solid", fgColor="17363D")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 26
    widths = [18, 18, 24, 18, 20, 18, 20, 18, 28, 16, 18, 20, 24, 20]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in range(2, TEMPLATE_LAST_ROW + 1):
        sheet.cell(row=row, column=8).number_format = "yyyy-mm-dd"

    options = workbook.create_sheet(_("选项"))
    option_headers = (
        "warehouse_code", _("仓库名称"),
        "category_code", _("分类名称"),
        "criticality", _("关键度"),
        "status", _("状态"),
    )
    options.append(option_headers)
    option_count = max(len(warehouses), len(categories), len(criticalities), len(statuses), 1)
    for index in range(option_count):
        warehouse = warehouses[index] if index < len(warehouses) else (None, None)
        category = categories[index] if index < len(categories) else (None, None)
        criticality = criticalities[index] if index < len(criticalities) else (None, None)
        status = statuses[index] if index < len(statuses) else (None, None)
        options.append([*warehouse, *category, *criticality, *status])

    _add_defined_name(workbook, "WarehouseCodes", "A", len(warehouses))
    _add_defined_name(workbook, "CategoryCodes", "C", len(categories))
    _add_defined_name(workbook, "CriticalityCodes", "E", len(criticalities))
    _add_defined_name(workbook, "EquipmentStatuses", "G", len(statuses))
    _add_list_validation(
        sheet,
        f"A2:A{TEMPLATE_LAST_ROW}",
        "WarehouseCodes",
        _("选择设备所属仓库编码。"),
    )
    _add_list_validation(
        sheet,
        f"D2:D{TEMPLATE_LAST_ROW}",
        "CategoryCodes",
        _("选择已存在的设备分类编码。"),
    )
    _add_list_validation(
        sheet,
        f"J2:J{TEMPLATE_LAST_ROW}",
        "CriticalityCodes",
        _("选择设备关键度。"),
    )
    _add_list_validation(
        sheet,
        f"K2:K{TEMPLATE_LAST_ROW}",
        "EquipmentStatuses",
        _("选择设备运行状态。"),
    )
    options.sheet_state = "hidden"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
