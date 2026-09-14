from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "MMS_需求规格与数据库设计_v1.0.docx"

FONT_LATIN = "Calibri"
FONT_CJK = "Microsoft YaHei"
BLUE = "1F5D7A"
DARK_BLUE = "183B4E"
LIGHT_BLUE = "E8F1F5"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "68737D"
WHITE = "FFFFFF"
INK = "1C2833"
GREEN = "2E6B57"
AMBER = "9A6700"

PAGE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 120, bottom: int = 90, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths: list[int], indent: int = TABLE_INDENT_DXA) -> None:
    if sum(widths) != PAGE_WIDTH_DXA:
        raise ValueError(f"Table widths must total {PAGE_WIDTH_DXA}, got {sum(widths)}")

    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(PAGE_WIDTH_DXA))

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(indent))

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[idx]))
            cell.width = Inches(widths[idx] / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run_font(run, size: float | None = None, color: str | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = FONT_LATIN
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 11.5, DARK_BLUE, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.10

    caption = styles["Caption"]
    caption.font.name = FONT_LATIN
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    caption.font.size = Pt(9)
    caption.font.color.rgb = RGBColor.from_string(MID_GRAY)
    caption.font.italic = False
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_with_next = True


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=8.5, color=MID_GRAY)


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)

    header = section.header
    p = header.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    run = p.add_run("MMS | 需求与数据模型规格\tV1.0")
    set_run_font(run, size=8.5, color=MID_GRAY, bold=True)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run("内部设计文档  |  第 ")
    set_run_font(run, size=8.5, color=MID_GRAY)
    add_page_field(p)
    run = p.add_run(" 页")
    set_run_font(run, size=8.5, color=MID_GRAY)


def add_title_block(doc: Document) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(34)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(8)
    run = kicker.add_run("PRODUCT REQUIREMENTS / DATA MODEL")
    set_run_font(run, size=9, color=GREEN, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("仓库自动化设备\n维保管理系统")
    set_run_font(run, size=28, color=DARK_BLUE, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(26)
    run = subtitle.add_run("Maintenance Management System (MMS)\n需求规格与数据库对应关系")
    set_run_font(run, size=14, color=BLUE)

    metadata = [
        ("文档版本", "V1.0"),
        ("编制日期", "2026-08-18"),
        ("适用阶段", "MVP 需求确认、原型开发与数据库建模"),
        ("目标读者", "业务负责人、仓库自动化团队、产品与研发团队"),
        ("推荐技术栈", "Django 5.2 LTS / PostgreSQL / Celery + Redis / Bootstrap 5"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(f"{label}：")
        set_run_font(r, size=10.5, color=MID_GRAY, bold=True)
        r = p.add_run(value)
        set_run_font(r, size=10.5, color=INK)

    callout = add_table(
        doc,
        ["设计结论"],
        [["系统以仓库为数据隔离边界，以设备与部件为维保对象，以周期计划驱动工单，并将提醒定向发送给对应仓库的工程师。"]],
        [PAGE_WIDTH_DXA],
        header_fill=BLUE,
        body_fill=LIGHT_BLUE,
    )
    callout.rows[0].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_page_break()


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True


def add_paragraph(doc: Document, text: str, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)


def add_bullets(doc: Document, items: list[str], numbered: bool = False) -> None:
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        p = doc.add_paragraph(style=style)
        r = p.add_run(item)
        set_run_font(r)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text, style="Caption")
    p.paragraph_format.keep_with_next = True


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[int],
    header_fill: str = BLUE,
    body_fill: str | None = None,
    small: bool = False,
) -> object:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = 0
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(text)
        set_run_font(r, size=8.5 if small else 9, color=WHITE, bold=True)
    for row_index, row_data in enumerate(rows):
        cells = table.add_row().cells
        for idx, text in enumerate(row_data):
            cell = cells[idx]
            if body_fill:
                set_cell_shading(cell, body_fill)
            elif row_index % 2 == 1:
                set_cell_shading(cell, "F8FAFB")
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if idx == 0 and len(headers) <= 4:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            elif len(text) <= 10:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            set_run_font(r, size=8 if small else 8.5, color=INK)
    set_table_geometry(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_section_break(doc: Document) -> None:
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def build_document() -> None:
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    add_title_block(doc)

    add_heading(doc, "1. 文档目的与产品定位")
    add_paragraph(
        doc,
        "本文档定义仓库自动化设备维保管理系统（MMS）第一阶段的业务范围、角色权限、核心流程、数据库实体关系、技术架构与验收条件。它同时作为网页框架、数据迁移和后续迭代的设计基线。",
    )
    add_heading(doc, "1.1 业务目标", 2)
    add_bullets(doc, [
        "建立跨仓库统一、可追溯的自动化设备与部件台账。",
        "允许按设备或具体部件设定日、周、月、季度、年等周期性维保计划。",
        "在到期前自动生成预防性维保工单，并通知设备所属仓库的工程师。",
        "记录工单接单、执行、检查、附件、完成和关闭全过程，形成审计链路。",
        "通过导入模板快速初始化存量设备，逐行反馈错误且不污染正式数据。",
    ])
    add_heading(doc, "1.2 成功指标（MVP）", 2)
    add_table(doc, ["指标", "目标值", "计算口径"], [
        ["设备台账覆盖率", ">= 95%", "已录入有效设备 / 盘点设备总数"],
        ["计划按期生成率", ">= 99%", "在计划触发窗口内成功生成的工单占比"],
        ["通知到达率", ">= 98%", "站内通知成功写入且具备接收人的占比"],
        ["工单可追溯率", "100%", "状态变化、执行人、时间与操作日志齐全"],
        ["导入错误可定位率", "100%", "错误包含批次、行号、字段与原因"],
    ], [2500, 1500, 5360])

    add_heading(doc, "2. 范围与边界")
    add_heading(doc, "2.1 MVP 范围", 2)
    add_bullets(doc, [
        "仓库、用户、仓库成员与工程师职责维护。",
        "设备分类、设备、设备部件及关键属性维护。",
        "CSV/XLSX 设备批量导入、预校验、错误报告与导入历史。",
        "维保模板、模板任务、维保计划、提前提醒天数和默认指派策略。",
        "周期扫描、幂等生成工单、站内通知和邮件扩展接口。",
        "工单列表、详情、接单、开始、完成、关闭、评论和附件。",
        "按仓库、状态、到期时间、设备、负责人筛选的基础看板。",
        "关键数据变更与业务操作审计。",
    ])
    add_heading(doc, "2.2 暂不纳入 MVP", 2)
    add_bullets(doc, [
        "备品备件库存、采购、供应商结算与成本核算。",
        "IoT/PLC 实时遥测、预测性维护模型和自动停机联动。",
        "原生移动 App、短信/企业微信/Teams 等多渠道正式集成。",
        "多级审批、电子签名、法规证书与外包服务商门户。",
        "离线作业、二维码打印与复杂排班优化。",
    ])

    add_heading(doc, "3. 角色、权限与数据隔离")
    add_caption(doc, "表 1 角色权限矩阵")
    add_table(doc, ["角色", "数据范围", "关键权限", "限制"], [
        ["平台管理员", "全部仓库", "用户、仓库、全局字典、审计与系统配置", "原则上不执行现场工单"],
        ["仓库管理员", "被授权仓库", "成员、设备、计划、工单分派、报表", "不可访问未授权仓库"],
        ["自动化工程师", "被授权仓库", "查看设备、接单、执行任务、提交结果", "不可删除基础台账或关闭他人工单"],
        ["维保主管", "被授权仓库", "计划审核、工单指派、验收关闭、逾期追踪", "不可维护平台级配置"],
        ["审计/只读", "被授权仓库", "查看台账、工单与审计日志", "无新增、修改、执行权限"],
    ], [1600, 1800, 3560, 2400], small=True)
    add_heading(doc, "3.1 权限规则", 2)
    add_bullets(doc, [
        "所有仓库业务表必须包含或可追溯到 warehouse_id；服务层查询默认按当前用户仓库成员关系过滤。",
        "平台管理员可跨仓库访问；其他用户必须通过有效 WarehouseMembership 获得访问权限。",
        "自动分派时优先使用计划默认负责人；未指定时通知该仓库所有启用通知的工程师和维保主管。",
        "任何 URL 参数中的仓库、设备或工单 ID 都必须再次进行对象级权限校验，不能只依赖前端隐藏。",
    ])

    add_heading(doc, "4. 功能需求")
    add_heading(doc, "4.1 仓库与人员", 2)
    add_bullets(doc, [
        "仓库具有唯一编码、名称、地址、时区、状态和负责人信息。",
        "一个用户可属于多个仓库，并可在不同仓库拥有不同角色。",
        "成员可设置主要仓库、通知开关和启用状态；离职或停用后不得接收新工单。",
    ])
    add_heading(doc, "4.2 设备与部件台账", 2)
    add_bullets(doc, [
        "每台设备必须归属且仅归属一个仓库，设备编码在仓库内唯一。",
        "设备支持分类、制造商、型号、序列号、投产日期、位置、关键度与状态。",
        "设备部件支持父子层级，例如输送机 > 驱动单元 > 电机；维保计划可指向整机或具体部件。",
        "设备停用后不再生成新周期工单，但历史工单和审计数据仍保留。",
    ])
    add_heading(doc, "4.3 设备导入", 2)
    add_bullets(doc, [
        "支持 CSV 和 XLSX，用户必须先选择目标仓库；导入文件不得通过字段覆盖目标仓库。",
        "先执行预校验再提交：必填、编码格式、枚举、日期、分类存在性、仓库内重复和文件内重复。",
        "默认采用 Upsert：以 warehouse_id + asset_code 定位；新增或更新均写入批次统计与审计日志。",
        "单行失败不阻断其他有效行；错误报告包含行号、字段、错误代码、消息和原始行数据。",
        "导入接口设置文件类型、大小、最大行数和公式单元格安全限制。",
    ])
    add_heading(doc, "4.4 维保模板与计划", 2)
    add_bullets(doc, [
        "维保模板定义可复用的标准任务，任务包括顺序、说明、预计工时和是否必须上传照片。",
        "维保计划关联仓库、设备、可选部件和模板，并配置周期数值、周期单位、首次到期日、提前天数、优先级与默认负责人。",
        "MVP 支持 DAY/WEEK/MONTH/QUARTER/YEAR；后续可增加运行小时、计数器或条件触发。",
        "同一设备或部件可有多个计划，例如每日外观检查、每月润滑、每季度电气紧固。",
        "修改周期时必须明确重新计算 next_due_at；暂停计划不生成新工单。",
    ])
    add_heading(doc, "4.5 自动生成工单与通知", 2)
    add_bullets(doc, [
        "调度任务定期查找 next_due_at - lead_days 已进入触发窗口且状态为 ACTIVE 的计划。",
        "生成工单时复制模板任务为工单任务快照，避免后续模板变化影响已发布工单。",
        "使用 source_plan_id + plan_due_at 唯一约束保证重试不重复生成。",
        "工单生成成功后推进计划 last_generated_at 与 next_due_at，并写入审计日志。",
        "通知收件人取默认负责人；若为空，则取工单仓库内启用通知的工程师与维保主管。",
        "站内通知为 MVP 必选渠道；邮件通过独立发送记录异步执行，失败可重试且不回滚工单。",
    ])
    add_heading(doc, "4.6 工单执行", 2)
    add_bullets(doc, [
        "工单状态：DRAFT、OPEN、ASSIGNED、IN_PROGRESS、COMPLETED、CLOSED、CANCELLED。",
        "工程师可接单、开始、逐项记录结果、上传附件和提交完成；主管负责关闭或退回。",
        "完成时校验必填任务、必传照片和完成说明；记录 started_at、completed_at、closed_at。",
        "状态变化必须追加 WorkOrderLog，不覆盖历史；取消必须填写原因。",
        "逾期定义为 due_at 早于当前时间且状态未进入 COMPLETED/CLOSED/CANCELLED。",
    ])

    add_heading(doc, "5. 关键业务流程")
    add_heading(doc, "5.1 设备导入流程", 2)
    add_bullets(doc, [
        "下载标准模板并选择目标仓库。",
        "上传文件，系统创建 ImportBatch 并完成格式与逐行校验。",
        "展示预计新增、更新、失败数量；用户确认后执行事务化分批写入。",
        "保存 ImportError 与汇总结果，提供错误明细下载。",
        "成功数据可进入设备详情，所有变更关联导入批次和操作者。",
    ], numbered=True)
    add_heading(doc, "5.2 周期计划到工单", 2)
    add_bullets(doc, [
        "调度器按仓库时区识别进入触发窗口的有效计划。",
        "数据库事务内锁定计划并检查 generation_key，防止并发重复。",
        "创建工单、任务快照和状态日志，然后计算下一到期时间。",
        "根据分派规则创建 NotificationRecipient，并异步发送各渠道。",
        "工程师在站内待办查看并执行，主管验收关闭，计划历史保持可追溯。",
    ], numbered=True)

    add_heading(doc, "6. 状态与业务规则")
    add_caption(doc, "表 2 工单状态转换")
    add_table(doc, ["当前状态", "允许动作", "目标状态", "操作者/校验"], [
        ["DRAFT", "发布", "OPEN / ASSIGNED", "主管；有负责人则直接 ASSIGNED"],
        ["OPEN", "分派 / 接单", "ASSIGNED", "主管或仓库工程师"],
        ["ASSIGNED", "开始执行", "IN_PROGRESS", "当前负责人"],
        ["IN_PROGRESS", "提交完成", "COMPLETED", "任务与附件规则全部通过"],
        ["COMPLETED", "验收关闭", "CLOSED", "维保主管或仓库管理员"],
        ["COMPLETED", "退回", "IN_PROGRESS", "必须填写退回原因"],
        ["DRAFT/OPEN/ASSIGNED", "取消", "CANCELLED", "主管；必须填写取消原因"],
    ], [1500, 1900, 1700, 4260], small=True)
    add_heading(doc, "6.1 核心不变量", 2)
    add_bullets(doc, [
        "计划、设备、部件、模板与默认负责人必须属于同一仓库上下文。",
        "部件 equipment_id 必须等于计划或工单 equipment_id；数据库约束不足部分由服务层校验。",
        "工单编号全局唯一，建议格式 WO-YYYYMMDD-######；编号生成需并发安全。",
        "已生成工单保存 plan_due_at 与任务快照；删除计划不得级联删除历史工单。",
        "业务记录优先软停用，不物理删除；附件按业务保留策略清理。",
        "所有时间以 UTC 入库，按仓库时区展示和计算日历周期。",
    ])

    add_heading(doc, "7. 数据库设计总览")
    add_paragraph(doc, "推荐生产数据库为 PostgreSQL 16+。本地框架默认使用 SQLite 以降低启动门槛，并通过环境变量切换 PostgreSQL。主键统一使用 UUID；金额外的常规计数使用整数；可扩展属性使用受控 JSONField。")
    add_heading(doc, "7.1 实体关系", 2)
    add_bullets(doc, [
        "User 与 Warehouse 通过 WarehouseMembership 构成多对多，并在关联上保存角色与通知设置。",
        "Warehouse 1:N Equipment；Equipment 1:N EquipmentComponent；Component 可通过 parent_id 自关联形成层级。",
        "MaintenanceTemplate 1:N MaintenanceTemplateTask；MaintenancePlan N:1 Template / Equipment，并可选关联 Component。",
        "MaintenancePlan 1:N WorkOrder；WorkOrder 1:N WorkOrderTask / WorkOrderLog。",
        "Notification 1:N NotificationRecipient；每条接收记录绑定一个 User 和一个渠道。",
        "ImportBatch 1:N ImportError；导入批次绑定 Warehouse 与创建人。",
        "Attachment 与 AuditLog 使用 entity_type + entity_id 关联不同业务对象，避免每类对象重复建表。",
    ])
    add_heading(doc, "7.2 业务模块与数据表对应", 2)
    add_table(doc, ["模块", "Django 应用 / 数据表", "主要职责"], [
        ["账号权限", "accounts.User", "登录身份、个人信息、平台管理员标志"],
        ["仓库组织", "facilities.Warehouse / WarehouseMembership", "仓库主数据、用户仓库角色和通知范围"],
        ["设备台账", "assets.AssetCategory / Equipment / EquipmentComponent", "设备分类、设备和部件层级"],
        ["维保标准", "maintenance.MaintenanceTemplate / MaintenanceTemplateTask", "可复用作业标准及任务步骤"],
        ["周期计划", "maintenance.MaintenancePlan", "周期、下次到期日、提前量和默认分派"],
        ["工单执行", "workorders.WorkOrder / WorkOrderTask / WorkOrderLog", "工单快照、执行任务和不可变操作轨迹"],
        ["消息中心", "notifications.Notification / NotificationRecipient", "消息内容、收件人、渠道与送达状态"],
        ["数据导入", "data_imports.ImportBatch / ImportError", "导入批次、统计、错误定位与原始数据"],
        ["通用支撑", "common.Attachment / AuditLog", "附件元数据与通用审计"],
    ], [1700, 3600, 4060], small=True)

    add_heading(doc, "8. 核心数据字典")
    add_paragraph(doc, "除特别说明外，业务表均包含 UUID 类型 id、created_at 和 updated_at。外键默认使用 RESTRICT/PROTECT 保护主数据；纯从属快照可使用 CASCADE。")

    dictionaries = [
        ("8.1 accounts_user", [
            ["username / email", "varchar", "唯一；至少一个可登录", "登录标识"],
            ["display_name", "varchar(100)", "必填", "姓名/显示名"],
            ["phone", "varchar(32)", "可空", "通知联系方式"],
            ["timezone", "varchar(64)", "默认 UTC", "个人展示时区"],
            ["is_active / is_staff", "boolean", "默认 true / false", "启用与后台权限"],
        ]),
        ("8.2 facilities_warehouse", [
            ["code", "varchar(32)", "唯一、不可变", "仓库编码"],
            ["name", "varchar(120)", "必填", "仓库名称"],
            ["address", "text", "可空", "物理地址"],
            ["timezone", "varchar(64)", "必填", "计划日历计算时区"],
            ["status", "enum", "ACTIVE/INACTIVE", "仓库状态"],
        ]),
        ("8.3 facilities_warehousemembership", [
            ["warehouse_id + user_id", "uuid FK", "联合唯一", "成员所属关系"],
            ["role", "enum", "ADMIN/ENGINEER/SUPERVISOR/VIEWER", "仓库内角色"],
            ["is_primary", "boolean", "每用户至多一个", "默认进入仓库"],
            ["notification_enabled", "boolean", "默认 true", "是否接收仓库通知"],
            ["is_active", "boolean", "默认 true", "成员关系是否生效"],
        ]),
        ("8.4 assets_equipment", [
            ["warehouse_id + asset_code", "uuid FK + varchar", "联合唯一", "仓库内设备标识"],
            ["name / category_id", "varchar + uuid FK", "必填", "设备名称与分类"],
            ["manufacturer / model / serial_number", "varchar", "可空", "厂家与铭牌数据"],
            ["location_detail / criticality", "varchar / enum", "可空 / 默认 MEDIUM", "位置与关键度"],
            ["commissioned_on / status", "date / enum", "可空 / ACTIVE", "投产日期与状态"],
            ["metadata", "jsonb", "默认 {}", "受控扩展属性"],
        ]),
        ("8.5 assets_equipmentcomponent", [
            ["equipment_id + code", "uuid FK + varchar", "联合唯一", "设备内部件编码"],
            ["parent_id", "uuid FK self", "可空；禁止环", "部件层级"],
            ["name / component_type", "varchar", "必填 / 可空", "名称与类型"],
            ["serial_number", "varchar", "可空", "部件序列号"],
            ["status", "enum", "ACTIVE/INACTIVE", "部件状态"],
        ]),
        ("8.6 maintenance_maintenancetemplate", [
            ["code + version", "varchar + integer", "联合唯一", "模板版本标识"],
            ["name / category_id", "varchar + uuid FK", "必填 / 可空", "模板名称与适用分类"],
            ["description", "text", "可空", "作业目标和说明"],
            ["is_active", "boolean", "默认 true", "是否可用于新计划"],
        ]),
        ("8.7 maintenance_maintenancetemplatetask", [
            ["template_id + sequence", "uuid FK + integer", "联合唯一", "模板内顺序"],
            ["title / instructions", "varchar / text", "必填 / 可空", "任务标题与操作说明"],
            ["estimated_minutes", "positive integer", "默认 0", "预计用时"],
            ["requires_photo", "boolean", "默认 false", "是否强制照片"],
        ]),
        ("8.8 maintenance_maintenanceplan", [
            ["warehouse_id / equipment_id", "uuid FK", "必填且一致", "计划所属仓库与设备"],
            ["component_id / template_id", "uuid FK", "部件可空；模板必填", "维保对象与作业标准"],
            ["cycle_unit / cycle_value", "enum / positive integer", "必填", "周期单位与数值"],
            ["start_date / next_due_at", "date / timestamptz", "必填", "锚点与下一到期时间"],
            ["lead_days / priority", "integer / enum", ">=0 / 默认 MEDIUM", "提前量与工单优先级"],
            ["default_assignee_id", "uuid FK", "可空且须为仓库成员", "默认负责人"],
            ["status / last_generated_at", "enum / timestamptz", "ACTIVE/PAUSED/ARCHIVED", "调度状态"],
        ]),
        ("8.9 workorders_workorder", [
            ["number", "varchar(32)", "全局唯一", "工单编号"],
            ["source_plan_id + plan_due_at", "uuid FK + timestamptz", "计划工单联合唯一", "幂等生成键"],
            ["warehouse/equipment/component_id", "uuid FK", "组件可空；上下文一致", "工单维保对象"],
            ["type / priority / status", "enum", "必填", "工单类型、优先级、状态"],
            ["title / description", "varchar / text", "必填 / 可空", "工单内容快照"],
            ["due_at / assignee_id", "timestamptz / uuid FK", "必填 / 可空", "期限与负责人"],
            ["started/completed/closed_at", "timestamptz", "按状态写入", "生命周期时间"],
        ]),
        ("8.10 workorders_workordertask / workorderlog", [
            ["work_order_id + sequence", "uuid FK + integer", "任务联合唯一", "工单任务顺序"],
            ["title / instructions / requires_photo", "多类型", "模板快照", "执行要求"],
            ["status / result", "enum / text", "必填 / 可空", "任务结果"],
            ["log.actor_id / action", "uuid FK / varchar", "操作者可空（系统）", "日志主体与动作"],
            ["from_status / to_status / note", "enum / text", "按动作填写", "不可变状态轨迹"],
        ]),
        ("8.11 notifications_notification / recipient", [
            ["type / title / message", "enum / varchar / text", "必填", "通知内容"],
            ["entity_type + entity_id", "varchar + uuid", "可空；联合索引", "关联业务对象"],
            ["dedup_key", "varchar(160)", "可空唯一", "消息幂等键"],
            ["notification_id + user_id + channel", "FK + FK + enum", "联合唯一", "接收人与渠道"],
            ["status / sent_at / read_at", "enum / timestamptz", "按投递更新", "送达及已读状态"],
            ["error_message / retry_count", "text / integer", "默认空 / 0", "发送失败诊断"],
        ]),
        ("8.12 data_imports_importbatch / importerror", [
            ["warehouse_id / created_by_id", "uuid FK", "必填", "导入范围与操作者"],
            ["source_file / import_type", "file / enum", "必填", "原文件与导入类型"],
            ["status", "enum", "UPLOADED/VALIDATING/READY/RUNNING/COMPLETED/FAILED", "批次状态"],
            ["total/success/failed_rows", "integer", ">=0", "批次统计"],
            ["row_number / field_name", "integer / varchar", "错误行必填", "错误位置"],
            ["error_code/message/raw_data", "varchar/text/jsonb", "必填", "错误诊断与原始行"],
        ]),
        ("8.13 common_attachment / auditlog", [
            ["entity_type + entity_id", "varchar + uuid", "联合索引", "通用对象引用"],
            ["file / original_name / mime_type / size", "多类型", "必填", "附件存储元数据"],
            ["uploaded_by_id", "uuid FK", "可空（系统）", "上传者"],
            ["audit.actor_id / warehouse_id", "uuid FK", "可空", "操作者与数据范围"],
            ["action / before_data / after_data", "varchar / jsonb", "动作必填", "审计差异"],
            ["ip_address / request_id", "inet / varchar", "可空", "请求追踪"],
        ]),
    ]
    for heading, rows in dictionaries:
        add_heading(doc, heading, 2)
        add_table(doc, ["字段", "类型", "约束/默认", "业务含义"], rows, [2400, 1760, 2400, 2800], small=True)

    add_heading(doc, "9. 约束、索引与一致性策略")
    add_caption(doc, "表 3 关键数据库约束与索引")
    add_table(doc, ["对象", "约束/索引", "目的"], [
        ["Warehouse", "UNIQUE(code)", "稳定的仓库业务键"],
        ["WarehouseMembership", "UNIQUE(warehouse_id, user_id)", "防止重复成员关系"],
        ["Equipment", "UNIQUE(warehouse_id, asset_code); INDEX(status, criticality)", "仓库内唯一与列表筛选"],
        ["EquipmentComponent", "UNIQUE(equipment_id, code); INDEX(parent_id)", "部件层级和快速定位"],
        ["MaintenancePlan", "INDEX(status, next_due_at); CHECK(cycle_value > 0, lead_days >= 0)", "高效扫描到期计划"],
        ["WorkOrder", "UNIQUE(source_plan_id, plan_due_at); INDEX(warehouse_id, status, due_at)", "生成幂等与待办查询"],
        ["NotificationRecipient", "UNIQUE(notification_id, user_id, channel); INDEX(user_id, status)", "防重复和消息中心查询"],
        ["ImportError", "INDEX(batch_id, row_number)", "按批次定位错误"],
        ["AuditLog", "INDEX(warehouse_id, created_at); INDEX(entity_type, entity_id)", "审计检索"],
    ], [2000, 4060, 3300], small=True)
    add_heading(doc, "9.1 事务边界", 2)
    add_bullets(doc, [
        "计划生成工单：锁定计划、检查幂等键、创建工单与任务、推进 next_due_at 在同一事务完成；通知在事务提交后投递。",
        "设备导入：预校验不写正式表；确认后分批事务写入，失败行保存错误，批次统计最终一致。",
        "工单状态转换：使用服务函数和条件更新，避免两个请求同时推进同一工单。",
        "审计日志与业务变更在同一事务内写入；高吞吐场景可使用 outbox 模式异步扩展。",
    ])

    add_heading(doc, "10. 技术架构与运行设计")
    add_table(doc, ["层次", "推荐组件", "用途"], [
        ["Web 与后台", "Django 5.2 LTS", "ORM、认证、权限、表单、Admin、服务端页面"],
        ["生产数据库", "PostgreSQL 16+", "事务、约束、JSONB、索引和并发控制"],
        ["本地数据库", "SQLite", "零配置启动与演示；不作为多实例生产方案"],
        ["异步任务", "Celery + Redis", "周期扫描、邮件发送、导入处理和重试"],
        ["页面样式", "Bootstrap 5 + 自定义 CSS", "响应式管理界面与一致组件"],
        ["文件存储", "本地开发 / S3 兼容对象存储", "导入文件与工单附件"],
        ["部署", "Gunicorn/Uvicorn + Nginx + Docker", "生产服务、静态文件和反向代理"],
    ], [1700, 2600, 5060])
    add_heading(doc, "10.1 Web 页面框架", 2)
    add_bullets(doc, [
        "登录页与仓库上下文选择。",
        "仪表盘：到期、逾期、进行中、近 7 日计划和仓库负载。",
        "设备：列表、详情、部件树、导入入口与导入历史。",
        "维保：模板、计划列表、计划创建/编辑和手动生成。",
        "工单：我的工单、仓库工单、详情、状态操作和执行记录。",
        "通知中心：未读计数、消息列表、已读状态。",
        "系统管理：仓库、成员、用户、字典、审计（按权限显示）。",
    ])

    add_heading(doc, "11. 非功能、安全与运维要求")
    add_bullets(doc, [
        "安全：CSRF、会话安全、密码策略、登录限速、文件白名单、对象级权限和最小权限。",
        "性能：普通列表 P95 < 2 秒；分页默认 25 条；到期扫描使用组合索引且按批处理。",
        "可靠性：调度和通知任务可重试、具备幂等键；重要任务暴露失败监控。",
        "审计：关键增删改、导入、分派和状态转换保留操作者、时间、对象与前后值。",
        "备份：PostgreSQL 每日备份并定期恢复演练；附件存储启用版本或生命周期策略。",
        "可观测性：结构化日志包含 request_id、warehouse_id、user_id；统计失败任务与逾期工单。",
        "兼容性：桌面 Chrome/Edge 最近两个主版本；核心操作适配手机浏览器。",
        "国际化：第一版中文，数据模型预留时区与后续多语言能力。",
    ])

    add_heading(doc, "12. MVP 验收标准")
    add_table(doc, ["编号", "验收场景", "通过条件"], [
        ["AC-01", "多仓库数据隔离", "工程师只能查看其有效成员关系对应仓库的数据"],
        ["AC-02", "设备导入", "合法行新增/更新成功；非法行展示准确行号、字段和原因"],
        ["AC-03", "部件级计划", "可为设备部件创建周期计划并正确展示下次到期时间"],
        ["AC-04", "自动工单", "进入触发窗口后生成且仅生成一张对应计划期次的工单"],
        ["AC-05", "定向通知", "默认负责人或所属仓库工程师收到站内消息，其他仓库不收到"],
        ["AC-06", "任务快照", "修改模板后，已生成工单的任务内容保持不变"],
        ["AC-07", "工单闭环", "接单、开始、完成、验收关闭均受权限和状态机约束"],
        ["AC-08", "逾期识别", "过期且未完成的工单在列表和仪表盘标识为逾期"],
        ["AC-09", "审计追踪", "导入、计划修改、工单状态变化可按对象查询完整日志"],
        ["AC-10", "基础可用性", "README 步骤可完成安装、迁移、创建管理员并启动系统"],
    ], [900, 2800, 5660], small=True)

    add_heading(doc, "13. 实施阶段建议")
    add_table(doc, ["阶段", "主要交付", "退出条件"], [
        ["阶段 1：框架", "项目配置、账号、仓库、设备、计划、工单模型与基础页面", "迁移通过，核心页面可访问"],
        ["阶段 2：业务闭环", "导入、调度生成、站内通知、工单状态机", "AC-01 至 AC-08 通过"],
        ["阶段 3：强化", "邮件、附件、审计查询、报表、PostgreSQL 与 Celery 部署", "可靠性与运维验收通过"],
        ["阶段 4：扩展", "IoT 计量触发、备件、移动端、外部消息渠道", "按新增业务目标评审"],
    ], [1700, 4260, 3400])

    add_heading(doc, "附录 A：设备导入模板")
    add_caption(doc, "表 4 equipment_import.xlsx 建议列")
    add_table(doc, ["列名", "必填", "示例", "校验/说明"], [
        ["asset_code", "是", "CV-001", "仓库内唯一；字母数字及 -_"],
        ["name", "是", "1 号主输送机", "1-120 字符"],
        ["category_code", "是", "CONVEYOR", "必须匹配有效设备分类"],
        ["manufacturer", "否", "Interroll", "最长 120 字符"],
        ["model", "否", "MCP-01", "最长 120 字符"],
        ["serial_number", "否", "SN202608001", "按业务可设仓库内唯一"],
        ["location_detail", "否", "A 区 2 层东侧", "最长 255 字符"],
        ["criticality", "否", "HIGH", "LOW/MEDIUM/HIGH/CRITICAL"],
        ["commissioned_on", "否", "2026-08-01", "ISO 日期 YYYY-MM-DD"],
        ["status", "否", "ACTIVE", "ACTIVE/INACTIVE/RETIRED"],
        ["component_code", "否", "MOTOR-01", "填写时同时创建/更新部件"],
        ["component_name", "条件必填", "驱动电机", "component_code 存在时必填"],
        ["component_type", "否", "MOTOR", "受控字典或自由文本（待确认）"],
    ], [1900, 1200, 2100, 4160], small=True)
    add_paragraph(doc, "导入时由用户在页面选择目标仓库，因此模板不包含 warehouse_code，防止跨仓库误写。部件层级的复杂导入建议在第二版独立提供 component_import.xlsx。")

    add_heading(doc, "附录 B：待业务确认事项")
    add_bullets(doc, [
        "工单默认提前生成天数，以及不同关键度设备是否需要不同默认值。",
        "通知渠道优先级：仅站内，还是第一版即启用企业邮箱。",
        "主管验收是否为所有工单必需，或仅 HIGH/CRITICAL 设备必需。",
        "设备编码能否更新；本设计建议一旦创建即视为稳定业务键。",
        "导入更新策略是否允许空值覆盖已有数据；本设计建议空值默认不覆盖。",
        "部件类型是否使用全局字典，以及是否需要针对设备分类限制可选部件。",
        "数据保留期、附件大小上限和审计日志保留年限。",
    ])

    doc.core_properties.title = "仓库自动化设备维保管理系统 - 需求规格与数据库设计"
    doc.core_properties.subject = "Maintenance Management System requirements and data model"
    doc.core_properties.author = "MMS Project Team"
    doc.core_properties.keywords = "MMS, maintenance, warehouse, automation, Django, PostgreSQL"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print("DOCX_CREATED")


if __name__ == "__main__":
    build_document()
