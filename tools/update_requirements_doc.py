from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


SOURCE = Path("docs/MMS_需求规格与数据库设计_v1.0.docx")


PARAGRAPH_REPLACEMENTS = {
    "文档版本：V1.0": "文档版本：V1.1",
    "编制日期：2026-08-18": "编制日期：2026-08-24",
    "适用阶段：MVP 需求确认、原型开发与数据库建模": "适用阶段：MVP 需求确认、框架开发、数据迁移与验收",
    "推荐技术栈：Django 5.2 LTS / PostgreSQL / Celery + Redis / Bootstrap 5": "推荐技术栈：Django 5.2 LTS / PostgreSQL / Celery + Redis / 自定义响应式 CSS",
    "允许按设备或具体部件设定日、周、月、季度、年等周期性维保计划。": "允许直接创建一次性工单或按设备、具体部件创建日、周、月、季度、年周期性工单。",
    "在到期前自动生成预防性维保工单，并通知设备所属仓库的工程师。": "周期源工单到期后自动生成独立执行工单，并通知设备所属仓库的可接单人员。",
    "维保模板、模板任务、维保计划、提前提醒天数和默认指派策略。": "一次性工单、周期性工单、工单任务、周期规则与负责人指派策略。",
    "周期扫描、幂等生成工单、站内通知和邮件扩展接口。": "周期工单扫描、幂等生成后续工单、站内通知和邮件扩展接口。",
    "工单列表、详情、接单、开始、完成、关闭、评论和附件。": "工单创建、列表、详情、接单、开始、完成、关闭、任务结果和附件。",
    "自动分派时优先使用计划默认负责人；未指定时通知该仓库所有启用通知的工程师和维保主管。": "创建工单时可直接选择负责人；未指定时通知该仓库所有角色允许接单且启用通知的成员。",
    "一个用户可属于多个仓库，并可在不同仓库拥有不同角色。": "一个用户可属于多个仓库，并可在不同仓库拥有不同的可配置角色；角色可在管理后台新增和修改。",
    "成员可设置主要仓库、通知开关和启用状态；离职或停用后不得接收新工单。": "成员可设置主要仓库、角色、通知开关和启用状态；离职或停用后不得接收新工单。",
    "设备部件支持父子层级，例如输送机 > 驱动单元 > 电机；维保计划可指向整机或具体部件。": "设备部件支持父子层级，例如输送机 > 驱动单元 > 电机；工单可指向整机或具体部件。",
    "设备停用后不再生成新周期工单，但历史工单和审计数据仍保留。": "设备停用后不可用于新工单，已有周期源不再生成后续工单；历史工单和审计数据仍保留。",
    "4.4 维保模板与计划": "4.4 一次性与周期性工单",
    "维保模板定义可复用的标准任务，任务包括顺序、说明、预计工时和是否必须上传照片。": "用户直接选择仓库、设备、可选部件、工单类型、优先级、到期时间和负责人创建工单，不再经过维保模板或维保计划。",
    "维保计划关联仓库、设备、可选部件和模板，并配置周期数值、周期单位、首次到期日、提前天数、优先级与默认负责人。": "一次性工单只执行一次；周期性工单额外配置周期单位、周期数值、可选结束时间、下次生成时间和启用状态。",
    "MVP 支持 DAY/WEEK/MONTH/QUARTER/YEAR；后续可增加运行小时、计数器或条件触发。": "周期单位支持 DAY/WEEK/MONTH/QUARTER/YEAR；后续可增加运行小时、计数器或条件触发。",
    "同一设备或部件可有多个计划，例如每日外观检查、每月润滑、每季度电气紧固。": "同一设备或部件可有多张周期工单，例如每日外观检查、每月润滑、每季度电气紧固。",
    "修改周期时必须明确重新计算 next_due_at；暂停计划不生成新工单。": "修改周期时重新计算 next_occurrence_at；停用或取消周期源后不再生成后续工单。",
    "调度任务定期查找 next_due_at - lead_days 已进入触发窗口且状态为 ACTIVE 的计划。": "调度任务定期查找 next_occurrence_at 已到期、recurrence_active 为 true 且未取消的周期源工单。",
    "生成工单时复制模板任务为工单任务快照，避免后续模板变化影响已发布工单。": "生成后续工单时复制周期源上的任务为独立任务快照，避免后续修改影响已发布工单。",
    "使用 source_plan_id + plan_due_at 唯一约束保证重试不重复生成。": "使用 recurrence_source_id + due_at 唯一约束保证重试不重复生成。",
    "工单生成成功后推进计划 last_generated_at 与 next_due_at，并写入审计日志。": "后续工单生成成功后推进周期源的 next_occurrence_at；超过结束时间时自动停用周期生成。",
    "通知收件人取默认负责人；若为空，则取工单仓库内启用通知的工程师与维保主管。": "通知收件人优先取工单负责人；若为空，则取工单仓库内角色允许接单且启用通知的成员。",
    "5.2 周期计划到工单": "5.2 周期工单生成后续工单",
    "调度器按仓库时区识别进入触发窗口的有效计划。": "调度器识别 next_occurrence_at 已到期的有效周期源工单。",
    "数据库事务内锁定计划并检查 generation_key，防止并发重复。": "数据库事务内锁定周期源，并检查 recurrence_source_id + due_at，防止并发重复。",
    "创建工单、任务快照和状态日志，然后计算下一到期时间。": "创建独立的一次性工单、任务快照和状态日志，然后计算下一生成时间。",
    "工程师在站内待办查看并执行，主管验收关闭，计划历史保持可追溯。": "工程师在站内待办查看并执行，主管验收关闭，周期源与每次执行历史保持可追溯。",
    "计划、设备、部件、模板与默认负责人必须属于同一仓库上下文。": "工单、设备、部件与负责人必须属于同一仓库上下文；下拉选项按仓库和设备逐级过滤。",
    "部件 equipment_id 必须等于计划或工单 equipment_id；数据库约束不足部分由服务层校验。": "部件 equipment_id 必须等于工单 equipment_id；数据库约束不足部分由模型、表单与服务层共同校验。",
    "已生成工单保存 plan_due_at 与任务快照；删除计划不得级联删除历史工单。": "自动生成工单保存 recurrence_source_id、due_at 与任务快照；取消周期源不得级联删除历史执行工单。",
    "User 与 Warehouse 通过 WarehouseMembership 构成多对多，并在关联上保存角色与通知设置。": "User 与 Warehouse 通过 WarehouseMembership 构成多对多；WarehouseRole 保存可配置角色、接单/管理标志和 Django 权限。",
    "MaintenanceTemplate 1:N MaintenanceTemplateTask；MaintenancePlan N:1 Template / Equipment，并可选关联 Component。": "历史 MaintenanceTemplate、MaintenanceTemplateTask 与 MaintenancePlan 表只读保留，用于数据兼容，不再承载新业务。",
    "MaintenancePlan 1:N WorkOrder；WorkOrder 1:N WorkOrderTask / WorkOrderLog。": "周期 WorkOrder 1:N 后续一次性 WorkOrder；WorkOrder 1:N WorkOrderTask / WorkOrderLog。",
    "8.3 facilities_warehousemembership": "8.3 facilities_warehouserole / warehousemembership",
    "8.6 maintenance_maintenancetemplate": "8.6 maintenance_maintenancetemplate（历史兼容）",
    "8.7 maintenance_maintenancetemplatetask": "8.7 maintenance_maintenancetemplatetask（历史兼容）",
    "8.8 maintenance_maintenanceplan": "8.8 maintenance_maintenanceplan（历史兼容）",
    "计划生成工单：锁定计划、检查幂等键、创建工单与任务、推进 next_due_at 在同一事务完成；通知在事务提交后投递。": "周期生成工单：锁定周期源、检查幂等键、创建工单与任务、推进 next_occurrence_at 在同一事务完成；通知随工单生成记录。",
    "仪表盘：到期、逾期、进行中、近 7 日计划和仓库负载。": "仪表盘：到期、逾期、进行中、近 7 日工单和周期工单队列。",
    "维保：模板、计划列表、计划创建/编辑和手动生成。": "工单创建：直接选择一次性或周期性执行，并联动选择仓库、设备、部件和负责人。",
    "系统管理：仓库、成员、用户、字典、审计（按权限显示）。": "系统管理：仓库、成员、可配置角色、用户、用户组、权限、字典和审计（按权限显示）。",
    "安全：CSRF、会话安全、密码策略、登录限速、文件白名单、对象级权限和最小权限。": "安全：CSRF、会话安全、登录限速、文件白名单、对象级权限和最小权限；当前不启用密码复杂度校验，管理员可设置简单密码。",
    "国际化：第一版中文，数据模型预留时区与后续多语言能力。": "国际化：网页与主要管理字段支持简体中文、英文和西班牙语，语言选择通过 Cookie 保持。",
    "工单默认提前生成天数，以及不同关键度设备是否需要不同默认值。": "周期工单默认周期、首次到期时间，以及不同关键度设备是否需要不同默认值。",
}


TABLE_REPLACEMENTS = {
    "系统以仓库为数据隔离边界，以设备与部件为维保对象，以周期计划驱动工单，并将提醒定向发送给对应仓库的工程师。": "系统以仓库为数据隔离边界，以设备与部件为维保对象，直接创建一次性或周期性工单，并将提醒定向发送给对应仓库的可接单人员。",
    "计划按期生成率": "周期工单按期生成率",
    "在计划触发窗口内成功生成的工单占比": "周期源到期后成功生成后续工单的占比",
    "成员、设备、计划、工单分派、报表": "成员、角色、设备、工单分派、报表",
    "计划审核、工单指派、验收关闭、逾期追踪": "周期工单管理、工单指派、验收关闭、逾期追踪",
    "仓库主数据、用户仓库角色和通知范围": "仓库主数据、可配置角色、用户成员关系和通知范围",
    "维保标准": "历史维保标准",
    "可复用作业标准及任务步骤": "只读保留的旧模板与任务，用于历史兼容",
    "周期计划": "历史维保计划",
    "周期、下次到期日、提前量和默认分派": "只读保留的旧计划；有效记录已迁移为周期工单",
    "工单快照、执行任务和不可变操作轨迹": "一次性/周期性工单、任务快照和不可变操作轨迹",
    "默认 true / false": "默认 true / false；加入用户组后自动启用后台资格",
    "启用与后台权限": "启用状态与后台登录资格；组权限直接生效",
    "是否可用于新计划": "历史字段；不再用于新业务",
    "模板版本标识": "历史模板版本标识",
    "模板名称与适用分类": "历史模板名称与适用分类",
    "作业目标和说明": "历史作业目标和说明",
    "模板内顺序": "历史模板内顺序",
    "预计用时": "历史预计用时",
    "计划所属仓库与设备": "历史计划所属仓库与设备",
    "维保对象与作业标准": "历史维保对象与作业标准",
    "周期单位与数值": "历史周期单位与数值",
    "锚点与下一到期时间": "历史锚点与下一到期时间",
    "提前量与工单优先级": "历史提前量与工单优先级",
    "默认负责人": "历史默认负责人",
    "调度状态": "历史调度状态",
    "source_plan_id + plan_due_at": "schedule_type / recurrence_active",
    "计划工单联合唯一": "ONE_TIME/RECURRING；默认一次性 / 默认 false",
    "幂等生成键": "工单频率与周期生成开关",
    "type / priority / status": "work_type / priority / status",
    "模板快照": "周期源任务快照",
    "MaintenancePlan": "Recurring WorkOrder",
    "INDEX(status, next_due_at); CHECK(cycle_value > 0, lead_days >= 0)": "INDEX(recurrence_active, next_occurrence_at); CHECK(recurrence_interval > 0)",
    "高效扫描到期计划": "高效扫描到期周期源",
    "UNIQUE(source_plan_id, plan_due_at); INDEX(warehouse_id, status, due_at)": "UNIQUE(recurrence_source_id, due_at); INDEX(warehouse_id, status, due_at)",
    "生成幂等与待办查询": "周期生成幂等与待办查询",
    "Bootstrap 5 + 自定义 CSS": "自定义响应式 CSS",
    "响应式管理界面与一致组件": "响应式业务界面与一致组件",
    "部件级计划": "部件级周期工单",
    "可为设备部件创建周期计划并正确展示下次到期时间": "可为设备部件直接创建周期工单并正确展示下次生成时间",
    "进入触发窗口后生成且仅生成一张对应计划期次的工单": "周期源到期后生成且仅生成一张对应期次的一次性工单",
    "修改模板后，已生成工单的任务内容保持不变": "修改周期源任务后，已生成工单的任务内容保持不变",
    "导入、计划修改、工单状态变化可按对象查询完整日志": "导入、周期工单修改、工单状态变化可按对象查询完整日志",
    "项目配置、账号、仓库、设备、计划、工单模型与基础页面": "项目配置、账号、仓库、设备、可配置角色、工单模型与基础页面",
}


def insert_paragraph_after(paragraph, text):
    new_element = OxmlElement("w:p")
    paragraph._p.addnext(new_element)
    new_paragraph = Paragraph(new_element, paragraph._parent)
    new_paragraph.style = "Normal"
    new_paragraph.add_run(text)
    return new_paragraph


def replace_paragraph_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def set_row(table, values):
    row = table.add_row()
    for cell, value in zip(row.cells, values):
        cell.text = value


def main():
    document = Document(SOURCE)

    seen_paragraphs = set()
    for paragraph in document.paragraphs:
        replacement = PARAGRAPH_REPLACEMENTS.get(paragraph.text)
        if replacement is not None:
            seen_paragraphs.add(paragraph.text)
            replace_paragraph_text(paragraph, replacement)

    missing_paragraphs = set(PARAGRAPH_REPLACEMENTS) - seen_paragraphs
    if missing_paragraphs:
        raise RuntimeError(f"Paragraph replacements not found: {sorted(missing_paragraphs)}")

    seen_cells = set()
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                replacement = TABLE_REPLACEMENTS.get(cell.text)
                if replacement is not None:
                    seen_cells.add(cell.text)
                    replace_paragraph_text(cell.paragraphs[0], replacement)
                    for paragraph in cell.paragraphs[1:]:
                        replace_paragraph_text(paragraph, "")

    missing_cells = set(TABLE_REPLACEMENTS) - seen_cells
    if missing_cells:
        raise RuntimeError(f"Table replacements not found: {sorted(missing_cells)}")

    membership_role_row = document.tables[7].rows[2]
    for cell, value in zip(
        membership_role_row.cells,
        (
            "role_id",
            "uuid FK",
            "引用可配置 WarehouseRole",
            "仓库内角色；角色可配置接单、管理和 Django 权限",
        ),
    ):
        replace_paragraph_text(cell.paragraphs[0], value)

    membership_heading = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "8.3 facilities_warehouserole / warehousemembership"
    )
    insert_paragraph_after(
        membership_heading,
        "WarehouseRole 通过 code、name、can_receive_work_orders、can_manage_work_orders、permissions 和 is_active 定义角色；WarehouseMembership 通过 role_id 引用角色。Django Group 权限无需复制到用户直接权限，分组后立即生效并自动启用后台登录资格。",
    )

    legacy_heading = next(
        paragraph
        for paragraph in document.paragraphs
        if paragraph.text == "8.6 maintenance_maintenancetemplate（历史兼容）"
    )
    insert_paragraph_after(
        legacy_heading,
        "以下三张 maintenance 表不再出现在业务页面或管理后台，仅为保留旧数据和旧工单外键而继续安装。迁移时，启用中的旧计划转换为周期性 WorkOrder。",
    )

    work_order_table = document.tables[13]
    set_row(
        work_order_table,
        (
            "recurrence_unit / recurrence_interval",
            "enum / positive integer",
            "周期工单必填；数值 > 0",
            "日、周、月、季度或年周期",
        ),
    )
    set_row(
        work_order_table,
        (
            "recurrence_end_at / next_occurrence_at",
            "timestamptz",
            "结束时间可空；下次生成可空",
            "周期边界与下一次生成时间",
        ),
    )
    set_row(
        work_order_table,
        (
            "recurrence_source_id + due_at",
            "uuid FK self + timestamptz",
            "自动生成工单联合唯一",
            "周期生成幂等键与执行期次",
        ),
    )

    acceptance_table = document.tables[20]
    set_row(
        acceptance_table,
        (
            "AC-11",
            "用户组权限",
            "用户加入组后无需配置直接权限即可获得组权限，并自动具备后台登录资格",
        ),
    )
    set_row(
        acceptance_table,
        (
            "AC-12",
            "多语言",
            "中文、英文、西班牙语可切换并保持选择，主要页面与表单同步翻译",
        ),
    )
    set_row(
        acceptance_table,
        (
            "AC-13",
            "简单密码",
            "用户创建表单允许密码 1 等简单密码，不执行复杂度校验",
        ),
    )

    output = SOURCE.with_suffix(".tmp.docx")
    document.core_properties.version = "1.1"
    document.core_properties.comments = "V1.1: simplified scheduling, configurable roles, group permissions, i18n, and simple passwords."
    document.save(output)
    output.replace(SOURCE)


if __name__ == "__main__":
    main()
