from datetime import timedelta
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.assets.models import AssetCategory, Equipment
from apps.facilities.models import Warehouse, WarehouseMembership, WarehouseRole
from apps.workorders.models import WorkOrder


class WorkOrderReportTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(
            username="report-user",
            email="report-user@example.com",
            password="1",
            timezone="America/Chicago",
        )
        role = WarehouseRole.objects.create(
            code="REPORT-VIEWER",
            name="报表查看人",
        )
        self.warehouse = Warehouse.objects.create(code="REPORT-A", name="报表一仓")
        other_warehouse = Warehouse.objects.create(code="REPORT-B", name="报表二仓")
        WarehouseMembership.objects.create(
            warehouse=self.warehouse,
            user=self.user,
            role=role,
        )
        category = AssetCategory.objects.create(code="REPORT-CAT", name="报表设备")
        equipment = Equipment.objects.create(
            warehouse=self.warehouse,
            category=category,
            asset_code="REPORT-EQ-A",
            name="可见设备",
        )
        other_equipment = Equipment.objects.create(
            warehouse=other_warehouse,
            category=category,
            asset_code="REPORT-EQ-B",
            name="不可见设备",
        )
        self.visible_order = WorkOrder.objects.create(
            number="WO-REPORT-VISIBLE",
            warehouse=self.warehouse,
            equipment=equipment,
            title="日报可见工单",
            due_at=timezone.now(),
        )
        WorkOrder.objects.create(
            number="WO-REPORT-HIDDEN",
            warehouse=other_warehouse,
            equipment=other_equipment,
            title="日报不可见工单",
            due_at=timezone.now(),
        )
        self.client.force_login(self.user)

    def test_daily_and_weekly_reports_are_valid_scoped_workbooks(self):
        for report_type in ("daily", "weekly"):
            with self.subTest(report_type=report_type):
                response = self.client.get(
                    reverse("dashboard:report", args=[report_type])
                )

                self.assertEqual(response.status_code, 200)
                self.assertEqual(
                    response["Content-Type"],
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
                workbook = load_workbook(BytesIO(response.content), data_only=False)
                self.assertEqual(workbook.sheetnames, ["汇总", "工单明细"])
                details = workbook["工单明细"]
                numbers = [details.cell(row=row, column=1).value for row in range(2, details.max_row + 1)]
                self.assertIn(self.visible_order.number, numbers)
                self.assertNotIn("WO-REPORT-HIDDEN", numbers)
                self.assertIsInstance(details["I2"].value, type(timezone.now().replace(tzinfo=None)))
                self.assertEqual(details["J2"].value, 24)
                self.assertIsInstance(details["K2"].value, type(timezone.now().replace(tzinfo=None)))
                self.assertEqual(details.auto_filter.ref, "A1:T2")
                self.assertEqual(details["T2"].value, self.visible_order.timezone)

    def test_dashboard_only_shows_started_orders_with_remaining_hours(self):
        WorkOrder.objects.create(
            number="WO-REPORT-FUTURE",
            warehouse=self.warehouse,
            equipment=self.visible_order.equipment,
            title="未来工单不应提前显示",
            due_at=timezone.now() + timedelta(days=1),
            task_duration_hours=4,
        )

        response = self.client.get(reverse("dashboard:home"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "日报可见工单")
        self.assertNotContains(response, "未来工单不应提前显示")
        self.assertContains(response, "24 小时")

    def test_reports_include_persisted_cycle_number_in_title(self):
        self.visible_order.recurrence_number = 3
        self.visible_order.save(update_fields=["recurrence_number"])
        response = self.client.get(reverse("dashboard:report", args=["daily"]))
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(workbook["工单明细"]["B2"].value, "日报可见工单 （第 3 期）")
