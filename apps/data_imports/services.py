import csv
import io
from datetime import date, datetime
from pathlib import Path

from django.db import transaction
from django.utils import timezone
from openpyxl import load_workbook

from apps.assets.models import AssetCategory, Equipment, EquipmentComponent

from .models import ImportBatch, ImportError


REQUIRED_HEADERS = {"asset_code", "name", "category_code"}


def _serializable(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return "" if value is None else str(value).strip()


def _read_rows(batch):
    batch.source_file.open("rb")
    file_obj = batch.source_file.file
    suffix = Path(batch.source_file.name).suffix.lower()
    if suffix == ".csv":
        text = io.TextIOWrapper(file_obj, encoding="utf-8-sig", newline="")
        yield from csv.DictReader(text)
        return
    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(iterator, [])]
    for values in iterator:
        yield dict(zip(headers, values, strict=False))


def _parse_date(value):
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.strptime(str(value).strip(), "%Y-%m-%d").date()


def _import_row(batch, raw_row):
    row = {str(key).strip(): _serializable(value) for key, value in raw_row.items() if key}
    missing_headers = REQUIRED_HEADERS - set(row)
    if missing_headers:
        raise ValueError(f"缺少必填列：{', '.join(sorted(missing_headers))}")
    if not row["asset_code"] or not row["name"] or not row["category_code"]:
        raise ValueError("asset_code、name 和 category_code 不能为空。")
    warehouse_code = row.get("warehouse_code", "")
    if warehouse_code and warehouse_code != batch.warehouse.code:
        raise ValueError(
            f"warehouse_code {warehouse_code} 与目标仓库 {batch.warehouse.code} 不一致。"
        )

    try:
        category = AssetCategory.objects.get(code=row["category_code"], is_active=True)
    except AssetCategory.DoesNotExist as exc:
        raise ValueError(f"设备分类 {row['category_code']} 不存在或已停用。") from exc

    criticality = row.get("criticality") or Equipment.Criticality.MEDIUM
    status = row.get("status") or Equipment.Status.ACTIVE
    if criticality not in Equipment.Criticality.values:
        raise ValueError(f"criticality 值无效：{criticality}")
    if status not in Equipment.Status.values:
        raise ValueError(f"status 值无效：{status}")

    equipment, _ = Equipment.objects.update_or_create(
        warehouse=batch.warehouse,
        asset_code=row["asset_code"],
        defaults={
            "name": row["name"],
            "category": category,
            "manufacturer": row.get("manufacturer", ""),
            "model": row.get("model", ""),
            "serial_number": row.get("serial_number", ""),
            "location_detail": row.get("location_detail", ""),
            "criticality": criticality,
            "commissioned_on": _parse_date(row.get("commissioned_on")),
            "status": status,
        },
    )
    component_code = row.get("component_code", "")
    if component_code:
        component_name = row.get("component_name", "")
        if not component_name:
            raise ValueError("填写 component_code 时必须同时填写 component_name。")
        EquipmentComponent.objects.update_or_create(
            equipment=equipment,
            code=component_code,
            defaults={
                "name": component_name,
                "component_type": row.get("component_type", ""),
            },
        )
    return row


def process_equipment_import(batch):
    batch.status = ImportBatch.Status.RUNNING
    batch.started_at = timezone.now()
    batch.save(update_fields=["status", "started_at", "updated_at"])
    batch.errors.all().delete()

    total = success = failed = 0
    try:
        rows = _read_rows(batch)
        for row_number, raw_row in enumerate(rows, start=2):
            if not any(value not in (None, "") for value in raw_row.values()):
                continue
            total += 1
            try:
                with transaction.atomic():
                    _import_row(batch, raw_row)
                success += 1
            except Exception as exc:
                failed += 1
                ImportError.objects.create(
                    batch=batch,
                    row_number=row_number,
                    error_code="ROW_VALIDATION_ERROR",
                    error_message=str(exc),
                    raw_data={str(key): _serializable(value) for key, value in raw_row.items() if key},
                )
        batch.status = ImportBatch.Status.COMPLETED
    except Exception as exc:
        batch.status = ImportBatch.Status.FAILED
        ImportError.objects.create(
            batch=batch,
            row_number=1,
            error_code="FILE_ERROR",
            error_message=str(exc),
        )
        failed += 1
    batch.total_rows = total
    batch.success_rows = success
    batch.failed_rows = failed
    batch.finished_at = timezone.now()
    batch.save(update_fields=["status", "total_rows", "success_rows", "failed_rows", "finished_at", "updated_at"])
    return batch
