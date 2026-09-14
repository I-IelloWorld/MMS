from collections import Counter
from datetime import datetime, time, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from django.db.models import Count, Q
from django.utils import timezone
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.common.timezones import DEFAULT_TIMEZONE
from apps.facilities.access import filter_by_warehouse
from apps.workorders.models import WorkOrder, WorkOrderTask


REPORT_TYPES = {"daily", "weekly"}
HEADER_FILL = PatternFill("solid", fgColor="17363D")
ACCENT_FILL = PatternFill("solid", fgColor="F1C84B")
SECTION_FILL = PatternFill("solid", fgColor="E9E6DE")
THIN_BORDER = Border(bottom=Side(style="thin", color="D8D3C9"))


def _user_zone(user):
    try:
        return ZoneInfo(user.timezone)
    except (AttributeError, ZoneInfoNotFoundError):
        return ZoneInfo(DEFAULT_TIMEZONE)


def report_window(report_type, user, now=None):
    if report_type not in REPORT_TYPES:
        raise ValueError("Unsupported report type")
    zone = _user_zone(user)
    local_now = timezone.localtime(now or timezone.now(), zone)
    report_date = local_now.date()
    if report_type == "weekly":
        report_date -= timedelta(days=report_date.weekday())
        end_date = report_date + timedelta(days=7)
    else:
        end_date = report_date + timedelta(days=1)
    start = datetime.combine(report_date, time.min, tzinfo=zone)
    end = datetime.combine(end_date, time.min, tzinfo=zone)
    return start, end, zone


def _orders_for_period(user, start, end):
    return list(
        filter_by_warehouse(
            WorkOrder.objects.select_related(
                "warehouse", "equipment", "component", "assignee", "created_by"
            ),
            user,
        )
        .filter(due_at__gte=start, due_at__lt=end)
        .annotate(
            task_count=Count("tasks", distinct=True),
            completed_task_count=Count(
                "tasks",
                filter=Q(tasks__status=WorkOrderTask.Status.DONE),
                distinct=True,
            ),
        )
        .order_by("due_at", "warehouse__code", "number")
    )


def _style_header(row):
    for cell in row:
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")


def _set_widths(sheet, widths):
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _excel_datetime(value, zone):
    if not value:
        return None
    return timezone.localtime(value, zone).replace(tzinfo=None)


def build_work_order_report(user, report_type, now=None):
    now = now or timezone.now()
    start, end, zone = report_window(report_type, user, now=now)
    orders = _orders_for_period(user, start, end)
    report_label = _("日报") if report_type == "daily" else _("周报")
    period_end = end - timedelta(microseconds=1)

    workbook = Workbook()
    summary = workbook.active
    summary.title = _("汇总")
    summary.sheet_view.showGridLines = False
    summary.merge_cells("A1:F1")
    summary["A1"] = _("MMS 工单%(report)s") % {"report": report_label}
    summary["A1"].fill = HEADER_FILL
    summary["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    summary["A1"].alignment = Alignment(vertical="center")
    summary.row_dimensions[1].height = 28
    summary["A2"] = _("统计周期")
    summary["B2"] = f"{start:%Y-%m-%d} - {period_end:%Y-%m-%d}"
    summary["A3"] = _("生成时间")
    summary["B3"] = _excel_datetime(now, zone)
    summary["B3"].number_format = "yyyy-mm-dd hh:mm"
    summary["D2"] = _("时区")
    summary["E2"] = str(zone)

    terminal = {
        WorkOrder.Status.COMPLETED,
        WorkOrder.Status.CLOSED,
        WorkOrder.Status.CANCELLED,
    }
    metrics = (
        (_("工单总数"), len(orders)),
        (_("待处理工单"), sum(order.status not in terminal for order in orders)),
        (_("已完成/关闭"), sum(order.status in {WorkOrder.Status.COMPLETED, WorkOrder.Status.CLOSED} for order in orders)),
        (_("已逾期"), sum(order.is_overdue for order in orders)),
    )
    for column, (label, value) in enumerate(metrics, start=1):
        cell = summary.cell(row=5, column=column, value=label)
        cell.fill = SECTION_FILL
        cell.font = Font(bold=True)
        summary.cell(row=6, column=column, value=value).font = Font(bold=True, size=18)

    summary["A8"] = _("状态分布")
    summary["A8"].fill = ACCENT_FILL
    summary["A8"].font = Font(bold=True)
    status_counts = Counter(order.status for order in orders)
    summary.append([_("状态"), _("数量")])
    _style_header(summary[9])
    row_number = 10
    for value, label in WorkOrder.Status.choices:
        summary.cell(row=row_number, column=1, value=str(label))
        summary.cell(row=row_number, column=2, value=status_counts[value])
        row_number += 1

    warehouse_column = 4
    summary.cell(row=8, column=warehouse_column, value=_("仓库分布"))
    summary.cell(row=8, column=warehouse_column).fill = ACCENT_FILL
    summary.cell(row=8, column=warehouse_column).font = Font(bold=True)
    summary.cell(row=9, column=warehouse_column, value=_("仓库"))
    summary.cell(row=9, column=warehouse_column + 1, value=_("数量"))
    _style_header(summary[9][warehouse_column - 1:warehouse_column + 1])
    warehouse_counts = Counter(
        f"{order.warehouse.code} · {order.warehouse.name}" for order in orders
    )
    for row_number, (warehouse, count) in enumerate(sorted(warehouse_counts.items()), start=10):
        summary.cell(row=row_number, column=warehouse_column, value=warehouse)
        summary.cell(row=row_number, column=warehouse_column + 1, value=count)
    _set_widths(summary, [24, 18, 4, 32, 14, 4])

    details = workbook.create_sheet(_("工单明细"))
    details.sheet_view.showGridLines = False
    headers = [
        _("工单编号"), _("标题"), _("仓库"), _("设备"), _("部件"),
        _("工单类型"), _("优先级"), _("状态"), _("工单开始生成日期"),
        _("任务时间（小时）"), _("逾期时间"), _("负责人"), _("创建人"),
        _("开始时间"), _("完成时间"), _("关闭时间"), _("任务数"),
        _("已完成任务数"), _("说明"), _("工单时区"),
    ]
    details.append(headers)
    _style_header(details[1])
    details.freeze_panes = "A2"
    details.auto_filter.ref = f"A1:T{max(1, len(orders) + 1)}"
    for order in orders:
        order_zone = ZoneInfo(order.timezone)
        details.append(
            [
                order.number,
                order.display_title,
                f"{order.warehouse.code} · {order.warehouse.name}",
                f"{order.equipment.asset_code} · {order.equipment.name}",
                order.component.name if order.component else _("整机"),
                order.get_work_type_display(),
                order.get_priority_display(),
                order.get_status_display(),
                _excel_datetime(order.due_at, order_zone),
                order.task_duration_hours,
                _excel_datetime(order.deadline_at, order_zone),
                str(order.assignee) if order.assignee else _("待认领"),
                str(order.created_by) if order.created_by else "",
                _excel_datetime(order.started_at, order_zone),
                _excel_datetime(order.completed_at, order_zone),
                _excel_datetime(order.closed_at, order_zone),
                order.task_count,
                order.completed_task_count,
                order.description,
                order.timezone,
            ]
        )
    for row in details.iter_rows(min_row=2):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in {2, 19})
        for column in (9, 11, 14, 15, 16):
            row[column - 1].number_format = "yyyy-mm-dd hh:mm"
    _set_widths(details, [22, 30, 28, 30, 20, 18, 12, 14, 20, 16, 20, 20, 20, 20, 20, 20, 10, 14, 42, 24])

    output = BytesIO()
    workbook.save(output)
    return output.getvalue(), start.date(), period_end.date()
